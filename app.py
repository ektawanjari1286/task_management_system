
from flask import (
    Flask, render_template, request, redirect,
    url_for, session, flash, abort, send_file
)
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime, timedelta, date
from functools import wraps
from getpass import getpass
from io import BytesIO
from xml.sax.saxutils import escape
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)

app = Flask(__name__)
app.config["SECRET_KEY"] = "CHANGE-THIS-TO-A-SECRET-RANDOM-KEY"
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///student_tasks.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)


# ===================== DATABASE MODELS =====================

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password = db.Column(db.String(200), nullable=False)
    role = db.Column(db.String(20), nullable=False, default="student")

    projects = db.relationship(
        "Project", backref="student", lazy=True,
        cascade="all, delete-orphan"
    )



class Project(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    subject = db.Column(db.String(100), nullable=False)
    topic = db.Column(db.String(200), nullable=False)
    deadline = db.Column(db.Date, nullable=False)
    # Soft delete: retain the project and its tasks for teacher/admin history.
    deleted_at = db.Column(db.DateTime, nullable=True, index=True)

    user_id = db.Column(
        db.Integer, db.ForeignKey("user.id"), nullable=False
    )

    tasks = db.relationship(
        "Task",
        backref="project",
        lazy=True,
        cascade="all, delete-orphan"
    )



class Task(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(300), nullable=False)
    description = db.Column(db.Text)
    due_date = db.Column(db.Date, nullable=False)
    status = db.Column(db.String(30), default="Pending")

    project_id = db.Column(
        db.Integer, db.ForeignKey("project.id"), nullable=False
    )


class Notification(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    message = db.Column(db.String(500), nullable=False)
    user_id = db.Column(
        db.Integer, db.ForeignKey("user.id"), nullable=False
    )
    is_read = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Feedback(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    message = db.Column(db.Text, nullable=False)

    project_id = db.Column(
        db.Integer, db.ForeignKey("project.id"), nullable=False
    )
    teacher_id = db.Column(
        db.Integer, db.ForeignKey("user.id"), nullable=False
    )
    student_id = db.Column(
        db.Integer, db.ForeignKey("user.id"), nullable=False
    )
    created_at = db.Column(
        db.DateTime, default=datetime.utcnow, nullable=False
    )

    project = db.relationship("Project", backref="feedback_entries")
    teacher = db.relationship("User", foreign_keys=[teacher_id])
    student = db.relationship("User", foreign_keys=[student_id])

# ===================== AUTHENTICATION =====================

def login_required(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            flash("Please login first.", "warning")
            return redirect(url_for("login"))
        return function(*args, **kwargs)
    return wrapper


def role_required(*roles):
    def decorator(function):
        @wraps(function)
        @login_required
        def wrapper(*args, **kwargs):
            user = db.session.get(User, session.get("user_id"))

            if not user:
                session.clear()
                flash("Please login again.", "warning")
                return redirect(url_for("login"))

            if user.role not in roles:
                flash(
                    "You do not have permission to access that page.",
                    "danger"
                )
                return redirect(url_for("home"))

            return function(*args, **kwargs)
        return wrapper
    return decorator


def redirect_to_role_dashboard(role):
    if role == "admin":
        return redirect(url_for("admin_dashboard"))
    if role == "teacher":
        return redirect(url_for("teacher_dashboard"))
    return redirect(url_for("dashboard"))


@app.route("/")
def home():
    if "user_id" in session:
        return redirect_to_role_dashboard(session.get("role"))
    return redirect(url_for("login"))


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        if not name or not email or not password:
            flash("Please fill in every field.", "danger")
            return redirect(url_for("register"))

        if len(password) < 8:
            flash("Password must be at least 8 characters.", "danger")
            return redirect(url_for("register"))

        if User.query.filter_by(email=email).first():
            flash("This email is already registered.", "danger")
            return redirect(url_for("register"))

        user = User(
            name=name,
            email=email,
            password=generate_password_hash(password),
            role="student"
        )
        db.session.add(user)
        db.session.commit()

        flash("Student registration successful. Please login.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect_to_role_dashboard(session.get("role"))

    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = User.query.filter_by(email=email).first()

        if user and check_password_hash(user.password, password):
            session.clear()
            session["user_id"] = user.id
            session["user_name"] = user.name
            session["role"] = user.role

            flash(f"Welcome, {user.name}!", "success")
            return redirect_to_role_dashboard(user.role)

        flash("Invalid email or password.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for("login"))


# ===================== DEADLINE ALERTS =====================

def create_deadline_notifications(user_id):
    """
    Creates each deadline reminder once, when the student opens
    the dashboard. Reminders are generated for the next 3 days.
    """
    today = date.today()
    reminder_limit = today + timedelta(days=3)
    projects = Project.query.filter_by(user_id=user_id, deleted_at=None).all()
    new_notifications = []

    for project in projects:
        for task in project.tasks:
            if task.status == "Completed":
                continue

            if task.due_date < today:
                message = (
                    f"OVERDUE ALERT [Task ID: {task.id}]: "
                    f"'{task.title}' was due on "
                    f"{task.due_date.strftime('%d %b %Y')}. "
                    "Please complete this task."
                )
            elif task.due_date <= reminder_limit:
                message = (
                    f"DEADLINE REMINDER [Task ID: {task.id}]: "
                    f"'{task.title}' is due on "
                    f"{task.due_date.strftime('%d %b %Y')}."
                )
            else:
                continue

            existing = Notification.query.filter_by(
                user_id=user_id,
                message=message
            ).first()

            if not existing:
                new_notifications.append(
                    Notification(user_id=user_id, message=message)
                )

    if new_notifications:
        db.session.add_all(new_notifications)
        db.session.commit()


# ===================== STUDENT DASHBOARD =====================


@app.route("/dashboard")
@role_required("student")
def dashboard():
    user_id = session["user_id"]

    # Generate deadline reminders and overdue alerts
    create_deadline_notifications(user_id)

    projects = Project.query.filter_by(
        user_id=user_id, deleted_at=None
    ).order_by(Project.id.desc()).all()

    notifications = Notification.query.filter_by(
        user_id=user_id,
        is_read=False
    ).order_by(Notification.created_at.desc()).all()

    all_tasks = [
        task
        for project in projects
        for task in project.tasks
    ]

    total_tasks = len(all_tasks)

    completed_tasks = sum(
        task.status == "Completed"
        for task in all_tasks
    )

    pending_tasks = total_tasks - completed_tasks

    overdue_tasks = sum(
        task.status != "Completed"
        and task.due_date < date.today()
        for task in all_tasks
    )

    progress = (
        round(completed_tasks * 100 / total_tasks)
        if total_tasks else 0
    )

    feedbacks = Feedback.query.filter_by(
        student_id=user_id
    ).order_by(Feedback.created_at.desc()).all()

    return render_template(
        "dashboard.html",
        projects=projects,
        notifications=notifications,
        total_tasks=total_tasks,
        completed_tasks=completed_tasks,
        pending_tasks=pending_tasks,
        overdue_tasks=overdue_tasks,
        progress=progress,
        feedbacks=feedbacks,
        today=date.today(),
        timedelta=timedelta
    )



# ===================== PROJECT CREATION =====================

@app.route("/create-project", methods=["GET", "POST"])
@role_required("student")
def create_project():
    if request.method == "POST":
        title = request.form.get("title", "").strip()
        subject = request.form.get("subject", "").strip()
        topic = request.form.get("topic", "").strip()
        deadline_text = request.form.get("deadline", "")

        if not title or not subject or not topic or not deadline_text:
            flash("Please fill in all project fields.", "danger")
            return redirect(url_for("create_project"))

        try:
            deadline = datetime.strptime(
                deadline_text, "%Y-%m-%d"
            ).date()
        except ValueError:
            flash("Please enter a valid deadline.", "danger")
            return redirect(url_for("create_project"))

        if deadline < date.today():
            flash("Deadline cannot be in the past.", "danger")
            return redirect(url_for("create_project"))

        project = Project(
            title=title,
            subject=subject,
            topic=topic,
            deadline=deadline,
            user_id=session["user_id"]
        )

        db.session.add(project)
        db.session.flush()

        generate_tasks(project)

        db.session.add(Notification(
            user_id=session["user_id"],
            message=f"Tasks have been generated for '{topic}'."
        ))
        db.session.commit()

        flash("Project created and task steps generated!", "success")
        return redirect(url_for("dashboard"))

    return render_template("create_project.html")


def generate_tasks(project):
    topic = project.topic.lower()

    if any(word in topic for word in [
        "machine learning", "classification", "prediction", "detection"
    ]):
        task_templates = [
            ("Define the problem and objectives",
             f"Define the problem, objectives and expected outcome for {project.topic}."),
            ("Collect and explore data",
             f"Find a suitable dataset for {project.topic} and inspect its columns and quality."),
            ("Prepare the dataset",
             "Clean missing values, handle duplicates and prepare features for modelling."),
            ("Build the model",
             f"Implement a suitable baseline model for {project.topic} using Python."),
            ("Evaluate model performance",
             "Evaluate the model with appropriate metrics and record the results."),
            ("Test and improve the solution",
             "Test with unseen examples and document limitations and possible improvements."),
            ("Prepare final documentation",
             f"Document methodology, results and conclusions for {project.topic}.")
        ]
    elif any(word in topic for word in [
        "database", "dbms", "reservation system"
    ]):
        task_templates = [
            ("Gather system requirements",
             f"Identify users, requirements and expected functions for {project.topic}."),
            ("Identify entities and relationships",
             "List entities, attributes, primary keys and relationships."),
            ("Design the ER diagram",
             "Draw the ER diagram and verify entity relationships."),
            ("Create database tables",
             "Create normalized tables with appropriate constraints and keys."),
            ("Implement SQL operations",
             "Write and test the required INSERT, UPDATE, DELETE and SELECT queries."),
            ("Validate database operations",
             "Test valid and invalid inputs, relationships and query results."),
            ("Prepare project documentation",
             "Document the ER diagram, schema, queries and test results.")
        ]
    elif any(word in topic for word in [
        "web", "website", "management system", "application"
    ]):
        task_templates = [
            ("Gather requirements",
             f"Define users, key features and requirements for {project.topic}."),
            ("Design the interface",
             "Create page layouts, navigation and wireframes for the required screens."),
            ("Build the frontend",
             "Implement the page structure and styling with suitable HTML and CSS."),
            ("Implement backend routes",
             "Build the required server routes, validation and business logic."),
            ("Connect the database",
             "Create or connect the required models and persist application data."),
            ("Test the application",
             "Test forms, navigation, access permissions and error handling."),
            ("Prepare final documentation",
             "Document features, architecture, screenshots and test results.")
        ]
    else:
        task_templates = [
            ("Define objectives",
             f"Define the problem, objectives and expected results for {project.topic}."),
            ("Research the topic",
             f"Study concepts, existing approaches and applications related to {project.topic}."),
            ("Plan the solution",
             "Break the project into components and prepare an implementation plan."),
            ("Develop the solution",
             f"Complete the main practical work related to {project.topic}."),
            ("Test and validate",
             "Test the solution against the project objectives and record results."),
            ("Improve the solution",
             "Resolve issues and make improvements based on test results."),
            ("Prepare documentation",
             f"Document objectives, methods, results and conclusions for {project.topic}.")
        ]

    today = date.today()
    days_available = max(0, (project.deadline - today).days)
    count = len(task_templates)

    for index, (title, description) in enumerate(task_templates):
        due_date = today + timedelta(
            days=round(days_available * (index + 1) / count)
        )
        due_date = min(due_date, project.deadline)

        db.session.add(Task(
            title=title,
            description=description,
            due_date=due_date,
            status="Pending",
            project_id=project.id
        ))


# ===================== TASK COMPLETION =====================

@app.route("/complete-task/<int:task_id>", methods=["GET", "POST"])
@role_required("student")
def complete_task(task_id):
    task = db.session.get(Task, task_id)

    if not task:
        abort(404)

    project = db.session.get(Project, task.project_id)

    if not project or project.user_id != session["user_id"]:
        abort(403)
    if project.deleted_at is not None:
        flash("This project is archived and can no longer be updated.", "warning")
        return redirect(url_for("dashboard"))

    if task.status != "Completed":
        task.status = "Completed"
        db.session.add(Notification(
            user_id=session["user_id"],
            message=f"Great job! You completed: {task.title}"
        ))
        db.session.commit()
        flash("Task completed successfully!", "success")
    else:
        flash("This task is already completed.", "info")

    return redirect(url_for("dashboard"))


@app.route("/notification-read/<int:notification_id>", methods=["GET", "POST"])
@role_required("student")
def notification_read(notification_id):
    notification = db.session.get(Notification, notification_id)

    if not notification:
        abort(404)

    if notification.user_id != session["user_id"]:
        abort(403)

    notification.is_read = True
    db.session.commit()
    return redirect(url_for("dashboard"))



# ===================== PROJECT DELETION =====================

@app.route("/delete-project/<int:project_id>", methods=["POST"])
@role_required("student")
def delete_project(project_id):
    try:
        project = db.session.get(Project, project_id)
        if project is None:
            flash("Project not found.", "error")
            return redirect(url_for("dashboard"))
        if project.user_id != session["user_id"]:
            abort(403)
        if project.deleted_at is not None:
            flash("This project is already in your archive.", "info")
            return redirect(url_for("dashboard"))

        # Archive rather than permanently deleting: teacher/admin can audit it.
        project.deleted_at = datetime.utcnow()
        db.session.add(Notification(
            user_id=session["user_id"],
            message=f"Project '{project.title}' was moved to your deleted-project history."
        ))
        db.session.commit()
        flash("Project moved to history. Your teacher and administrator can still view it.", "success")
    except Exception as e:
        db.session.rollback()
        app.logger.exception("Project archive failed: %s", e)
        flash("Project could not be archived. Check the terminal for details.", "error")
    return redirect(url_for("dashboard"))


# ===================== TEACHER DASHBOARD =====================

@app.route("/teacher")
@role_required("teacher")
def teacher_dashboard():
    students = User.query.filter_by(role="student").order_by(
        User.name.asc()
    ).all()

    student_data = []

    for student in students:
        project_details = []
        student_total_tasks = 0
        student_completed_tasks = 0

        for project in sorted(student.projects, key=lambda item: item.id, reverse=True):
            tasks = project.tasks
            completed = sum(
                1 for task in tasks if task.status == "Completed"
            )
            total = len(tasks)
            pending = total - completed
            progress = round(
                (completed / total) * 100
            ) if total else 0

            if project.deleted_at is None:
                student_total_tasks += total
                student_completed_tasks += completed

            project_feedback = Feedback.query.filter_by(
                project_id=project.id
            ).order_by(Feedback.created_at.desc()).all()

            project_details.append({
                "project": project,
                "is_deleted": project.deleted_at is not None,
                "tasks": tasks,
                "total": total,
                "completed": completed,
                "pending": pending,
                "progress": progress,
                "feedbacks": project_feedback
            })

        student_data.append({
            "student": student,
            "projects": project_details,
            "total_tasks": student_total_tasks,
            "completed_tasks": student_completed_tasks,
            "pending_tasks": student_total_tasks - student_completed_tasks
        })

    return render_template(
        "teacher.html",
        student_data=student_data,
        today=date.today(),
        timedelta=timedelta
    )


@app.route("/teacher/feedback/<int:project_id>", methods=["POST"])
@role_required("teacher")
def submit_feedback(project_id):
    project = db.session.get(Project, project_id)

    if not project:
        abort(404)
    if project.deleted_at is not None:
        flash("Archived projects are read-only; feedback cannot be added.", "warning")
        return redirect(url_for("teacher_dashboard"))

    message = request.form.get("message", "").strip()

    if not message:
        flash("Please enter feedback before submitting.", "warning")
        return redirect(url_for("teacher_dashboard"))

    if len(message) > 3000:
        flash("Feedback must be 3000 characters or fewer.", "warning")
        return redirect(url_for("teacher_dashboard"))

    feedback = Feedback(
        message=message,
        project_id=project.id,
        teacher_id=session["user_id"],
        student_id=project.user_id
    )

    db.session.add(feedback)
    db.session.add(Notification(
        user_id=project.user_id,
        message=f"Your teacher has added feedback for '{project.title}'."
    ))
    db.session.commit()

    flash("Feedback submitted successfully!", "success")
    return redirect(url_for("teacher_dashboard"))


# ===================== REPORT DATA =====================

def collect_report_data(students):
    today = date.today()

    summary_rows = []
    task_rows = []
    feedback_rows = []

    for student in students:
        projects = sorted(
            student.projects,
            key=lambda item: item.id
        )

        if not projects:
            summary_rows.append([
                student.name, student.email, "No projects",
                "", "", "", 0, 0, 0, 0, 0
            ])

        for project in projects:
            tasks = sorted(project.tasks, key=lambda item: item.due_date)

            total = len(tasks)
            completed = sum(
                task.status == "Completed" for task in tasks
            )
            pending = total - completed
            overdue = sum(
                task.status != "Completed" and task.due_date < today
                for task in tasks
            )
            progress = round(
                completed * 100 / total
            ) if total else 0

            summary_rows.append([
                student.name,
                student.email,
                project.title,
                project.subject,
                project.topic,
                project.deadline.strftime("%d-%b-%Y"),
                total,
                completed,
                pending,
                overdue,
                f"{progress}%"
            ])

            for task in tasks:
                is_overdue = (
                    task.status != "Completed"
                    and task.due_date < today
                )

                task_rows.append([
                    student.name,
                    project.title,
                    task.title,
                    task.description or "",
                    task.due_date.strftime("%d-%b-%Y"),
                    task.status,
                    "Yes" if is_overdue else "No"
                ])

            feedbacks = Feedback.query.filter_by(
                project_id=project.id
            ).order_by(Feedback.created_at.desc()).all()

            for feedback in feedbacks:
                feedback_rows.append([
                    student.name,
                    project.title,
                    feedback.teacher.name,
                    feedback.message,
                    feedback.created_at.strftime("%d-%b-%Y %I:%M %p")
                ])

    return summary_rows, task_rows, feedback_rows


# ===================== EXCEL REPORT GENERATOR =====================

def create_excel_report(students, report_title):
    summary_rows, task_rows, feedback_rows = collect_report_data(students)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Project Summary"

    summary_sheet.append([
        "Student Name", "Student Email", "Project", "Subject",
        "Topic", "Project Deadline", "Total Tasks", "Completed",
        "Pending", "Overdue", "Progress"
    ])

    for row in summary_rows:
        summary_sheet.append(row)

    task_sheet = workbook.create_sheet("Task Details")
    task_sheet.append([
        "Student Name", "Project", "Task", "Description",
        "Due Date", "Status", "Overdue"
    ])

    for row in task_rows:
        task_sheet.append(row)

    feedback_sheet = workbook.create_sheet("Teacher Feedback")
    feedback_sheet.append([
        "Student Name", "Project", "Teacher",
        "Feedback", "Date"
    ])

    for row in feedback_rows:
        feedback_sheet.append(row)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="18213E"
    )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = 0

            for cell in column_cells:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )
                value = str(cell.value or "")
                max_length = max(max_length, min(len(value), 45))

            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 3, 14), 45
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        "student_progress_report.xlsx"
        if report_title == "Student Progress Report"
        else "teacher_progress_report.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# ===================== PDF REPORT GENERATOR =====================

def create_pdf_report(students, report_title):
    summary_rows, task_rows, feedback_rows = collect_report_data(students)

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(escape(report_title), styles["Title"]))
    story.append(Paragraph(
        "Generated on: " + date.today().strftime("%d %B %Y"),
        styles["Normal"]
    ))
    story.append(Spacer(1, 7 * mm))

    story.append(Paragraph("Project Summary", styles["Heading2"]))

    summary_table_data = [[
        "Student", "Project", "Subject", "Deadline",
        "Tasks", "Done", "Pending", "Overdue", "Progress"
    ]]

    for row in summary_rows:
        summary_table_data.append([
            Paragraph(escape(str(row[0])), styles["BodyText"]),
            Paragraph(escape(str(row[2])), styles["BodyText"]),
            Paragraph(escape(str(row[3])), styles["BodyText"]),
            str(row[5]),
            str(row[6]),
            str(row[7]),
            str(row[8]),
            str(row[9]),
            str(row[10])
        ])

    summary_table = Table(
        summary_table_data,
        repeatRows=1,
        colWidths=[80, 125, 70, 65, 42, 38, 45, 45, 52]
    )
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18213E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DCE1EE")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white, colors.HexColor("#F3F5FB")
        ]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5)
    ]))

    story.append(summary_table)
    story.append(Spacer(1, 7 * mm))

    story.append(Paragraph("Task Details", styles["Heading2"]))

    task_table_data = [[
        "Student", "Project", "Task", "Due Date", "Status", "Overdue"
    ]]

    for row in task_rows:
        task_table_data.append([
            Paragraph(escape(str(row[0])), styles["BodyText"]),
            Paragraph(escape(str(row[1])), styles["BodyText"]),
            Paragraph(escape(str(row[2])), styles["BodyText"]),
            str(row[4]),
            str(row[5]),
            str(row[6])
        ])

    if not task_rows:
        task_table_data.append(["No tasks available", "", "", "", "", ""])

    task_table = Table(
        task_table_data,
        repeatRows=1,
        colWidths=[85, 115, 250, 65, 65, 55]
    )
    task_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18213E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DCE1EE")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white, colors.HexColor("#F3F5FB")
        ]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5)
    ]))

    story.append(task_table)
    story.append(Spacer(1, 7 * mm))

    story.append(Paragraph("Teacher Feedback", styles["Heading2"]))

    feedback_table_data = [[
        "Student", "Project", "Teacher", "Feedback", "Date"
    ]]

    for row in feedback_rows:
        feedback_table_data.append([
            Paragraph(escape(str(row[0])), styles["BodyText"]),
            Paragraph(escape(str(row[1])), styles["BodyText"]),
            Paragraph(escape(str(row[2])), styles["BodyText"]),
            Paragraph(escape(str(row[3])), styles["BodyText"]),
            str(row[4])
        ])

    if not feedback_rows:
        feedback_table_data.append([
            "No feedback available", "", "", "", ""
        ])

    feedback_table = Table(
        feedback_table_data,
        repeatRows=1,
        colWidths=[85, 110, 85, 285, 75]
    )
    feedback_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18213E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DCE1EE")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5)
    ]))

    story.append(feedback_table)
    document.build(story)
    output.seek(0)

    filename = (
        "student_progress_report.pdf"
        if report_title == "Student Progress Report"
        else "teacher_progress_report.pdf"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf"
    )


# ===================== REPORT DOWNLOAD ROUTES =====================

@app.route("/student/report/excel")
@role_required("student")
def student_report_excel():
    student = db.session.get(User, session["user_id"])
    return create_excel_report([student], "Student Progress Report")


@app.route("/student/report/pdf")
@role_required("student")
def student_report_pdf():
    student = db.session.get(User, session["user_id"])
    return create_pdf_report([student], "Student Progress Report")


@app.route("/teacher/report/excel")
@role_required("teacher")
def teacher_report_excel():
    students = User.query.filter_by(role="student").order_by(
        User.name.asc()
    ).all()
    return create_excel_report(students, "Teacher Progress Report")


@app.route("/teacher/report/pdf")
@role_required("teacher")
def teacher_report_pdf():
    students = User.query.filter_by(role="student").order_by(
        User.name.asc()
    ).all()
    return create_pdf_report(students, "Teacher Progress Report")


# ===================== ADMIN DASHBOARD =====================

@app.route("/admin")
@role_required("admin")
def admin_dashboard():
    users = User.query.order_by(User.id.desc()).all()

    return render_template(
        "admin.html",
        users=users,
        total_users=User.query.count(),
        total_students=User.query.filter_by(role="student").count(),
        total_teachers=User.query.filter_by(role="teacher").count(),
        total_projects=Project.query.filter_by(deleted_at=None).count(),
        deleted_projects=Project.query.filter(Project.deleted_at.isnot(None)).count(),
        projects=Project.query.order_by(Project.id.desc()).all(),
        total_tasks=Task.query.count(),
        completed_tasks=Task.query.filter_by(status="Completed").count()
    )


@app.route("/admin/create-user", methods=["POST"])
@role_required("admin")
def admin_create_user():
    name = request.form.get("name", "").strip()
    email = request.form.get("email", "").strip().lower()
    password = request.form.get("password", "")
    role = request.form.get("role", "")

    if not name or not email or not password:
        flash("All user fields are required.", "danger")
        return redirect(url_for("admin_dashboard"))

    if role not in ("teacher", "student"):
        flash("Admin can create student or teacher accounts only.", "danger")
        return redirect(url_for("admin_dashboard"))

    if len(password) < 8:
        flash("Password must be at least 8 characters.", "danger")
        return redirect(url_for("admin_dashboard"))

    if User.query.filter_by(email=email).first():
        flash("This email is already registered.", "danger")
        return redirect(url_for("admin_dashboard"))

    db.session.add(User(
        name=name,
        email=email,
        password=generate_password_hash(password),
        role=role
    ))
    db.session.commit()

    flash(f"{role.capitalize()} account created successfully.", "success")
    return redirect(url_for("admin_dashboard"))


# ===================== DATABASE / ADMIN SETUP =====================

with app.app_context():
    db.create_all()
    # Upgrade existing SQLite databases safely without requiring data deletion.
    from sqlalchemy import inspect, text
    inspector = inspect(db.engine)
    if "project" in inspector.get_table_names():
        columns = {column["name"] for column in inspector.get_columns("project")}
        if "deleted_at" not in columns:
            with db.engine.begin() as connection:
                connection.execute(text("ALTER TABLE project ADD COLUMN deleted_at DATETIME"))


def create_admin_account():
    with app.app_context():
        print("Create the first administrator account")
        name = input("Admin name: ").strip()
        email = input("Admin email: ").strip().lower()
        password = getpass("Admin password (minimum 8 characters): ")
        confirm = getpass("Confirm password: ")

        if not name or not email or len(password) < 8:
            print(
                "Name/email are required; password must have "
                "at least 8 characters."
            )
            return

        if password != confirm:
            print("Passwords do not match.")
            return

        if User.query.filter_by(email=email).first():
            print("This email already exists. No account was changed.")
            return

        db.session.add(User(
            name=name,
            email=email,
            password=generate_password_hash(password),
            role="admin"
        ))
        db.session.commit()
        print("Administrator account created successfully.")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "create-admin":
        create_admin_account()
    else:
        app.run(debug=True)
